import sys
import os
import re
import json
import sqlite3
import datetime
import pandas as pd
import duckdb
from docx import Document
from pathlib import Path
from typing import Dict, List, Any, Tuple, Optional

PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, PROJECT_ROOT)

from scripts.anonymize_private_data import detect_pii_in_dataframe

def log_operation(operation: str, details: dict):
    """將操作記錄寫入 logs/processing_log.jsonl (絕不記錄個資真實內容)"""
    log_dir = os.path.join(PROJECT_ROOT, "logs")
    os.makedirs(log_dir, exist_ok=True)
    log_file = os.path.join(log_dir, "processing_log.jsonl")
    
    log_entry = {
        "timestamp": datetime.datetime.now().isoformat(),
        "operation": operation,
        "details": details
    }
    
    with open(log_file, "a", encoding="utf-8") as f:
        f.write(json.dumps(log_entry, ensure_ascii=False) + "\n")

def sanitize_table_name(name: str) -> str:
    """轉換檔案名或工作表名為合法的 SQL 表格名稱"""
    # 僅保留英數字、中文、底線
    sanitized = re.sub(r'[^\w\u4e00-\u9fa5]', '_', name)
    # 確保不以數字開頭
    if sanitized and sanitized[0].isdigit():
        sanitized = "t_" + sanitized
    return sanitized

def import_excel(file_path: str) -> Dict[str, pd.DataFrame]:
    """匯入 Excel 檔案的所有工作表"""
    xls = pd.ExcelFile(file_path)
    sheets_data = {}
    for sheet_name in xls.sheet_names:
        df = pd.read_excel(file_path, sheet_name=sheet_name)
        # 清理欄位名稱 (防 SQL 報錯)
        df.columns = [str(col).strip() for col in df.columns]
        sheets_data[sheet_name] = df
    return sheets_data

def import_csv(file_path: str) -> Dict[str, pd.DataFrame]:
    """匯入 CSV 檔案"""
    # 嘗試不同編碼
    for encoding in ['utf-8', 'utf-8-sig', 'big5', 'gbk', 'latin1']:
        try:
            df = pd.read_csv(file_path, encoding=encoding)
            df.columns = [str(col).strip() for col in df.columns]
            return {"default": df}
        except UnicodeDecodeError:
            continue
    raise ValueError(f"無法讀取 CSV 檔案：編碼不支援")

def import_docx(file_path: str) -> Dict[str, pd.DataFrame]:
    """匯入 Word 檔案 (擷取文字段落及表格)"""
    doc = Document(file_path)
    
    # 1. 擷取段落文字
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    
    # 2. 擷取表格
    tables_data = []
    for t_idx, table in enumerate(doc.tables):
        rows = []
        for row in table.rows:
            rows.append([cell.text.strip() for cell in row.cells])
        if rows:
            headers = rows[0]
            # 處理重複或空的表頭
            seen_headers = {}
            cleaned_headers = []
            for i, h in enumerate(headers):
                if not h:
                    h = f"Column_{i}"
                if h in seen_headers:
                    seen_headers[h] += 1
                    h = f"{h}_{seen_headers[h]}"
                else:
                    seen_headers[h] = 0
                cleaned_headers.append(h)
                
            data = rows[1:] if len(rows) > 1 else []
            df = pd.DataFrame(data, columns=cleaned_headers)
            tables_data.append((f"Table_{t_idx}", df))
            
    # 將段落也包裝成一個簡單的 DataFrame
    text_df = pd.DataFrame({"paragraph_index": range(len(paragraphs)), "content": paragraphs})
    
    result = {"paragraphs": text_df}
    for name, df in tables_data:
        result[name] = df
    return result

def import_text(file_path: str) -> Dict[str, pd.DataFrame]:
    """匯入 TXT/MD 檔案"""
    with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
        lines = f.readlines()
    
    cleaned_lines = [line.strip() for line in lines if line.strip()]
    df = pd.DataFrame({"line_number": range(1, len(cleaned_lines) + 1), "content": cleaned_lines})
    return {"default": df}

def import_file(file_path: str, db_path: str = None) -> Tuple[bool, str]:
    """主匯入函數"""
    file_path = os.path.abspath(file_path)
    if not os.path.exists(file_path):
        return False, "檔案不存在"
        
    path_obj = Path(file_path)
    ext = path_obj.suffix.lower()
    file_name = path_obj.name
    
    if db_path is None:
        db_path = os.path.join(PROJECT_ROOT, "database", "local_index.duckdb")
    os.makedirs(os.path.dirname(db_path), exist_ok=True)
    
    try:
        # 依副檔名選取讀取方式
        if ext in ['.xlsx', '.xls']:
            data_dict = import_excel(file_path)
        elif ext == '.csv':
            data_dict = import_csv(file_path)
        elif ext == '.docx':
            data_dict = import_docx(file_path)
        elif ext in ['.txt', '.md']:
            data_dict = import_text(file_path)
        else:
            return False, f"不支援的檔案類型: {ext}"
            
        conn = duckdb.connect(db_path)
        
        # 初始化 file_summaries 表
        conn.execute('''
        CREATE TABLE IF NOT EXISTS file_summaries (
            file_name VARCHAR,
            file_path VARCHAR,
            file_type VARCHAR,
            sheet_name VARCHAR,
            table_name VARCHAR,
            column_names VARCHAR,
            row_count INTEGER,
            has_pii BOOLEAN,
            pii_types VARCHAR,
            imported_at TIMESTAMP,
            is_anonymized BOOLEAN,
            safe_for_cloud BOOLEAN,
            PRIMARY KEY (file_path, sheet_name)
        )
        ''')
        
        for sheet_name, df in data_dict.items():
            # 轉換資料欄位型態為字串，方便 DuckDB 儲存與後續模糊比對
            # 避免混合型態導致 DuckDB 寫入錯誤
            df_str = df.astype(str)
            
            # 進行 PII 評估
            pii_report = detect_pii_in_dataframe(df, file_name=f"{file_name} -> {sheet_name}")
            
            sanitized_file = sanitize_table_name(path_obj.stem)
            sanitized_sheet = sanitize_table_name(sheet_name)
            table_name = f"data_{sanitized_file}_{sanitized_sheet}"
            
            # 寫入資料表 (若已存在則覆蓋)
            conn.register('temp_df', df_str)
            conn.execute(f"DROP TABLE IF EXISTS {table_name}")
            conn.execute(f"CREATE TABLE {table_name} AS SELECT * FROM temp_df")
            conn.unregister('temp_df')
            
            # 更新摘要
            conn.execute(
                "DELETE FROM file_summaries WHERE file_path = ? AND sheet_name = ?",
                (file_path, sheet_name)
            )
            
            conn.execute(
                "INSERT INTO file_summaries VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
                (
                    file_name,
                    file_path,
                    ext,
                    sheet_name,
                    table_name,
                    ",".join(df.columns),
                    len(df),
                    len(pii_report.pii_columns) > 0,
                    ",".join(pii_report.pii_types_found),
                    datetime.datetime.now(),
                    False,
                    False
                )
            )
            
            # 寫入 logs/processing_log.jsonl
            log_operation("import_file", {
                "file_name": file_name,
                "file_type": ext,
                "sheet_name": sheet_name,
                "table_name": table_name,
                "row_count": len(df),
                "has_pii": len(pii_report.pii_columns) > 0,
                "pii_types": pii_report.pii_types_found
            })
            
        conn.close()
        return True, "匯入完成"
        
    except Exception as e:
        log_operation("import_error", {"file_path": file_path, "error": str(e)})
        return False, f"匯入失敗：{str(e)}"

def get_imported_files(db_path: str = None) -> List[Dict[str, Any]]:
    """獲取所有已匯入檔案的摘要"""
    if db_path is None:
        db_path = os.path.join(PROJECT_ROOT, "database", "local_index.duckdb")
    if not os.path.exists(db_path):
        return []
        
    conn = duckdb.connect(db_path)
    try:
        res = conn.execute("SELECT * FROM file_summaries").fetchall()
        cols = [desc[0] for desc in conn.description]
        results = []
        for r in res:
            results.append(dict(zip(cols, r)))
        conn.close()
        return results
    except Exception:
        conn.close()
        return []

if __name__ == '__main__':
    # 簡單測試
    print("匯入模組測試")
    test_path = os.path.join(PROJECT_ROOT, "tests", "test_excel.xlsx")
    # 如果有檔案可以測試
    if os.path.exists(test_path):
        success, msg = import_file(test_path)
        print(f"結果：{success}, 訊息：{msg}")
        summaries = get_imported_files()
        print(f"目前資料庫中檔案數：{len(summaries)}")
    else:
        print(f"無測試檔案：{test_path}")
